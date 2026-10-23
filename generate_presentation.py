import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions
    NAVY = RGBColor(15, 23, 42)       # #0F172A
    DEEP_BLUE = RGBColor(30, 58, 138)  # #1E3A8A
    BLUE = RGBColor(37, 99, 235)      # #2563EB
    LIGHT_BLUE = RGBColor(219, 234, 254) # #DBEAFE
    GREEN = RGBColor(16, 185, 129)    # #10B981
    LIGHT_GREEN = RGBColor(209, 250, 229) # #D1FAE5
    SLATE_DARK = RGBColor(30, 41, 59) # #1E293B
    SLATE_MUTED = RGBColor(100, 116, 139) # #64748B
    BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC
    WHITE = RGBColor(255, 255, 255)
    AMBER = RGBColor(245, 158, 11)    # #F59E0B
    LIGHT_AMBER = RGBColor(254, 243, 199) # #FEF3C7
    BORDER_COLOR = RGBColor(226, 232, 240) # #E2E8F0

    def add_header(slide, title_text, category_text="CAMPUSPLAY • DESIGN THINKING MINI-PROJECT"):
        # Header background banner
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = NAVY
        header_shape.line.fill.background()

        # Accent line below banner
        line_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.1), Inches(13.333), Inches(0.06))
        line_shape.fill.solid()
        line_shape.fill.fore_color.rgb = BLUE
        line_shape.line.fill.background()

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.28))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(9.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = GREEN

        # Title text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.5), Inches(0.55))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    def add_footer(slide, current_slide, total_slides=15):
        # Footer text
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(11.733), Inches(0.3))
        tf = footer_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Vidyalankar Institute of Technology (VIT) | Department of Information Technology | Slide {current_slide} of {total_slides}"
        p.font.size = Pt(8.5)
        p.font.color.rgb = SLATE_MUTED

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Visual Accent Bar
    bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.1), Inches(0.18), Inches(4.8))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = BLUE
    bar1.line.fill.background()

    # Title Box
    tbox1 = slide1.shapes.add_textbox(Inches(1.4), Inches(1.1), Inches(10.8), Inches(3.2))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    
    p0 = tf1.paragraphs[0]
    p0.text = "DESIGN THINKING MINI-PROJECT PRESENTATION"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = GREEN

    p1 = tf1.add_paragraph()
    p1.text = "CampusPlay 🎾"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_before = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Enhancing Campus Sports Access Through Automated Scheduling, Equipment Governance & Academic Timetable Conflict Verification"
    p2.font.size = Pt(17)
    p2.font.color.rgb = LIGHT_BLUE
    p2.space_before = Pt(8)

    # Info Card Box
    card_info = add_card(slide1, Inches(1.4), Inches(4.6), Inches(10.8), Inches(1.85), bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
    ibox = slide1.shapes.add_textbox(Inches(1.6), Inches(4.75), Inches(10.4), Inches(1.5))
    itf = ibox.text_frame
    itf.word_wrap = True

    ip1 = itf.paragraphs[0]
    ip1.text = "Institution: Vidyalankar Institute of Technology (VIT), Wadala, Mumbai"
    ip1.font.size = Pt(13)
    ip1.font.bold = True
    ip1.font.color.rgb = WHITE

    ip2 = itf.add_paragraph()
    ip2.text = "Department: Department of Information Technology | Academic Level: Second Year (S.E.) Engineering"
    ip2.font.size = Pt(12)
    ip2.font.color.rgb = LIGHT_GREEN
    ip2.space_before = Pt(4)

    ip3 = itf.add_paragraph()
    ip3.text = "Project Lead: Aryan Kelkar (Roll No: [TO BE CONFIRMED BY TEAM] | aryan.kelkar.2016@gmail.com)"
    ip3.font.size = Pt(12)
    ip3.font.color.rgb = WHITE
    ip3.space_before = Pt(4)

    ip4 = itf.add_paragraph()
    ip4.text = "Project Team: [TO BE CONFIRMED BY PROJECT TEAM: Member 2, Member 3, Member 4] | Faculty Mentor: [TO BE CONFIRMED BY TEAM]"
    ip4.font.size = Pt(11)
    ip4.font.color.rgb = SLATE_MUTED
    ip4.space_before = Pt(4)

    # ==========================================
    # SLIDE 2: PROJECT OVERVIEW
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Project Overview: Executive Summary")
    add_footer(slide2, 2)

    # 3 Cards Layout
    # Card 1: What is CampusPlay
    add_card(slide2, Inches(0.8), Inches(1.4), Inches(3.64), Inches(5.4))
    b1 = slide2.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(3.24), Inches(5.0))
    t1 = b1.text_frame
    t1.word_wrap = True
    p = t1.paragraphs[0]
    p.text = "🎯 What is CampusPlay?"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BLUE

    bullets1 = [
        ("Institutional Platform", "A comprehensive, full-stack campus sports facility management platform engineered specifically for collegiate sports infrastructure."),
        ("Multi-Facility Support", "Centralized scheduling for outdoor grounds (Ground 1, Ground 2) and indoor courts (Basketball Court, Badminton Court)."),
        ("Dual Role Architecture", "Dedicated role-based portals for collegiate students and sports department administrators."),
        ("Single Source of Truth", "Eliminates physical register books, verbal booking disputes, and uncoordinated ground allocations.")
    ]
    for title, desc in bullets1:
        p = t1.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(10)

    # Card 2: The Core Problem & Users
    add_card(slide2, Inches(4.84), Inches(1.4), Inches(3.64), Inches(5.4))
    b2 = slide2.shapes.add_textbox(Inches(5.04), Inches(1.6), Inches(3.24), Inches(5.0))
    t2 = b2.text_frame
    t2.word_wrap = True
    p = t2.paragraphs[0]
    p.text = "👥 Target Users & Problem"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE

    bullets2 = [
        ("Target Student Body", "Undergraduate students across engineering streams seeking recreational & competitive sporting access."),
        ("Sports Administrators", "Faculty coordinators and sports directors managing venue availability, asset allocation, and approvals."),
        ("Information Fragmentation", "Students previously had zero visibility into ground availability without physically visiting the sports gymkhana."),
        ("Academic Clash Hazard", "Ground reservations frequently coincided with mandatory academic lectures and laboratory practicals.")
    ]
    for title, desc in bullets2:
        p = t2.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(10)

    # Card 3: Value Proposition & Impact
    add_card(slide2, Inches(8.88), Inches(1.4), Inches(3.64), Inches(5.4))
    b3 = slide2.shapes.add_textbox(Inches(9.08), Inches(1.6), Inches(3.24), Inches(5.0))
    t3 = b3.text_frame
    t3.word_wrap = True
    p = t3.paragraphs[0]
    p.text = "⚡ Value & Innovation"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GREEN

    bullets3 = [
        ("Real-Time Slot Discovery", "Zero-friction availability grid reflecting vacant, booked, and reserved time-slots across 11 daily slots."),
        ("Academic Conflict Check", "Automated verification against institutional division timetables (INFT, CMPN, EXTC, etc.) preventing bunking."),
        ("Equipment Governance", "End-to-end checkout, return inspection, damage tracking, and inventory transaction auditing."),
        ("Instant Live Sync", "WebSocket-driven instant notifications for immediate approval/rejection visibility without refreshing.")
    ]
    for title, desc in bullets3:
        p = t3.add_paragraph()
        p.text = f"• {title}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(10)

    # ==========================================
    # SLIDE 3: PROBLEM STATEMENT
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Problem Statement: Challenges in Campus Sports Access")
    add_footer(slide3, 3)

    # 4 Problem Cards
    cards_data = [
        ("1. Opaque Availability & Slot Collisions", 
         "Students arrive at grounds only to discover existing sessions or unannounced closures. The lack of a digital availability tracker leads to friction, double-booking arguments, and wasted student recreation time.",
         AMBER, LIGHT_AMBER),
        ("2. Academic Schedule Desynchronization", 
         "Students frequently book recreational facilities during allocated lecture or lab hours without faculty awareness. Administrators lacked any mechanism to cross-reference academic timetables before approving slots.",
         RGBColor(239, 68, 68), RGBColor(254, 226, 226)),
        ("3. Unaccountable Sports Equipment Depletion", 
         "Sports gear (bats, footballs, racquets) was issued through paper logbooks. Equipment returned in damaged condition or never returned remained untracked, causing institutional budget depletion and shortages.",
         BLUE, LIGHT_BLUE),
        ("4. Administrative Overhead & Lack of Audit", 
         "Sports department staff spent excessive hours answering in-person inquiries, mediating scheduling conflicts, and maintaining physical paper logs vulnerable to loss, alteration, and zero searchability.",
         DEEP_BLUE, RGBColor(241, 245, 249))
    ]

    for idx, (title, desc, color, bg) in enumerate(cards_data):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.35 + row * 2.3)
        add_card(slide3, left, top, Inches(5.75), Inches(2.1), bg_color=bg, border_color=color)

        tb = slide3.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.35), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = SLATE_DARK
        p2.space_before = Pt(6)

    # Synthesized Problem Statement at Bottom
    add_card(slide3, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9), bg_color=NAVY, border_color=BLUE)
    tbs = slide3.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.3), Inches(0.8))
    tfs = tbs.text_frame
    tfs.word_wrap = True
    ps = tfs.paragraphs[0]
    ps.text = "SYNTHESIZED PROBLEM STATEMENT:"
    ps.font.size = Pt(10)
    ps.font.bold = True
    ps.font.color.rgb = GREEN

    ps2 = tfs.add_paragraph()
    ps2.text = "\"How might we provide collegiate students with transparent, real-time sports facility reservation while empowering sports administrators with automated academic timetable conflict detection and verifiable equipment governance?\""
    ps2.font.size = Pt(12)
    ps2.font.bold = True
    ps2.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 4: TARGET USERS AND EMPATHY
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Target Users & Empathy Mapping")
    add_footer(slide4, 4)

    # 3 Persona Columns
    # Persona 1: Competitive College Athlete
    add_card(slide4, Inches(0.8), Inches(1.4), Inches(3.64), Inches(5.4))
    b1 = slide4.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(3.24), Inches(5.1))
    t1 = b1.text_frame
    t1.word_wrap = True
    p = t1.paragraphs[0]
    p.text = "User Persona: Student Athlete"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BLUE

    p_sub = t1.add_paragraph()
    p_sub.text = "E.g., Rohan (S.E. IT, Football Captain)"
    p_sub.font.size = Pt(10)
    p_sub.font.italic = True
    p_sub.font.color.rgb = SLATE_MUTED

    p_g = t1.add_paragraph()
    p_g.text = "Goals & Needs:"
    p_g.font.bold = True
    p_g.font.size = Pt(11)
    p_g.font.color.rgb = SLATE_DARK
    p_g.space_before = Pt(8)
    p_g2 = t1.add_paragraph()
    p_g2.text = "• Book 2-hour practice sessions in advance for tournament preparation\n• Register team squad members with verified collegiate credentials\n• Access quality football gear reliably"
    p_g2.font.size = Pt(10.5)
    p_g2.font.color.rgb = SLATE_DARK

    p_p = t1.add_paragraph()
    p_p.text = "Frustrations & Pain Points:"
    p_p.font.bold = True
    p_p.font.size = Pt(11)
    p_p.font.color.rgb = RGBColor(239, 68, 68)
    p_p.space_before = Pt(8)
    p_p2 = t1.add_paragraph()
    p_p2.text = "• Travelling to campus only to find ground allocated to casual matches\n• Equipment missing or broken on match days\n• Bureaucratic delay in getting physical sign-offs"
    p_p2.font.size = Pt(10.5)
    p_p2.font.color.rgb = SLATE_DARK

    # Persona 2: Recreational Student
    add_card(slide4, Inches(4.84), Inches(1.4), Inches(3.64), Inches(5.4))
    b2 = slide4.shapes.add_textbox(Inches(5.04), Inches(1.55), Inches(3.24), Inches(5.1))
    t2 = b2.text_frame
    t2.word_wrap = True
    p = t2.paragraphs[0]
    p.text = "User Persona: Casual Player"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN

    p_sub2 = t2.add_paragraph()
    p_sub2.text = "E.g., Priya (T.E. CMPN, Badminton Enthusiast)"
    p_sub2.font.size = Pt(10)
    p_sub2.font.italic = True
    p_sub2.font.color.rgb = SLATE_MUTED

    p_g = t2.add_paragraph()
    p_g.text = "Goals & Needs:"
    p_g.font.bold = True
    p_g.font.size = Pt(11)
    p_g.font.color.rgb = SLATE_DARK
    p_g.space_before = Pt(8)
    p_g2 = t2.add_paragraph()
    p_g2.text = "• Quick 1-hour badminton slot during free academic intervals\n• Instant mobile verification and slot booking\n• Borrow racquets & shuttlecocks without complex forms"
    p_g2.font.size = Pt(10.5)
    p_g2.font.color.rgb = SLATE_DARK

    p_p = t2.add_paragraph()
    p_p.text = "Frustrations & Pain Points:"
    p_p.font.bold = True
    p_p.font.size = Pt(11)
    p_p.font.color.rgb = RGBColor(239, 68, 68)
    p_p.space_before = Pt(8)
    p_p2 = t2.add_paragraph()
    p_p2.text = "• Lack of transparency on whether court is vacant\n• Uncertainty about sports room operating hours\n• Fear of unintended attendance clash with labs"
    p_p2.font.size = Pt(10.5)
    p_p2.font.color.rgb = SLATE_DARK

    # Persona 3: Sports Administrator / Faculty
    add_card(slide4, Inches(8.88), Inches(1.4), Inches(3.64), Inches(5.4))
    b3 = slide4.shapes.add_textbox(Inches(9.08), Inches(1.55), Inches(3.24), Inches(5.1))
    t3 = b3.text_frame
    t3.word_wrap = True
    p = t3.paragraphs[0]
    p.text = "User Persona: Sports Admin"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE

    p_sub3 = t3.add_paragraph()
    p_sub3.text = "E.g., Sports Director / Gymkhana In-Charge"
    p_sub3.font.size = Pt(10)
    p_sub3.font.italic = True
    p_sub3.font.color.rgb = SLATE_MUTED

    p_g = t3.add_paragraph()
    p_g.text = "Goals & Needs:"
    p_g.font.bold = True
    p_g.font.size = Pt(11)
    p_g.font.color.rgb = SLATE_DARK
    p_g.space_before = Pt(8)
    p_g2 = t3.add_paragraph()
    p_g2.text = "• Verify student eligibility & timetable conflicts prior to approving\n• Ensure equipment is physically checked on return and reconciled\n• Maintain centralized digital log for institutional audits"
    p_g2.font.size = Pt(10.5)
    p_g2.font.color.rgb = SLATE_DARK

    p_p = t3.add_paragraph()
    p_p.text = "Frustrations & Pain Points:"
    p_p.font.bold = True
    p_p.font.size = Pt(11)
    p_p.font.color.rgb = RGBColor(239, 68, 68)
    p_p.space_before = Pt(8)
    p_p2 = t3.add_paragraph()
    p_p2.text = "• Endless paper registers with illegible handwriting\n• Students playing during lecture hours creating attendance disputes\n• Unreturned and damaged gear with zero student accountability"
    p_p2.font.size = Pt(10.5)
    p_p2.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 5: DESIGN THINKING PROCESS
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Design Thinking Methodology: 5-Stage Iterative Cycle")
    add_footer(slide5, 5)

    stages = [
        ("1. EMPATHIZE", "Understand User Realities", 
         "• Investigated campus sports access patterns\n• Observed gymkhana desk bottlenecks\n• Documented student frustrations with ground conflicts",
         "Stakeholder Pain Point Matrix", BLUE),
        ("2. DEFINE", "Synthesize Problem Core", 
         "• Formulated actionable Point-Of-View\n• Established constraints: college domain (@vit.edu.in), division rosters\n• Framed HMW questions",
         "Core System Requirements", DEEP_BLUE),
        ("3. IDEATE", "Explore Architectural Options", 
         "• Compared Form vs Custom Web App\n• Devised automated timetable conflict engine\n• Structured 2-step equipment return reconciliation",
         "Solution Blueprint & Feature Spec", GREEN),
        ("4. PROTOTYPE", "Build & Refine System", 
         "• Multi-step Booking Wizard\n• Sport & Nature Harmony UI\n• Admin verification panel & live WebSocket synchronization",
         "Full-Stack Web Application", AMBER),
        ("5. TEST", "Verify Rigorously", 
         "• 49 Automated unit/integration tests\n• Verified interval overlap edge cases\n• Validated file magic bytes & role access controls",
         "49/49 Verified Passing Tests", RGBColor(16, 185, 129))
    ]

    for idx, (stitle, ssub, sacts, sartifact, scolor) in enumerate(stages):
        left = Inches(0.8 + idx * 2.38)
        top = Inches(1.4)
        add_card(slide5, left, top, Inches(2.25), Inches(5.4), bg_color=WHITE, border_color=scolor)

        # Stage Banner
        sb = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(2.25), Inches(0.65))
        sb.fill.solid()
        sb.fill.fore_color.rgb = scolor
        sb.line.fill.background()
        sbtf = sb.text_frame
        sbtf.word_wrap = True
        sp = sbtf.paragraphs[0]
        sp.text = stitle
        sp.font.size = Pt(12)
        sp.font.bold = True
        sp.font.color.rgb = WHITE
        sp.alignment = PP_ALIGN.CENTER

        # Stage Content
        tb = slide5.shapes.add_textbox(left + Inches(0.12), top + Inches(0.75), Inches(2.01), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = ssub
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = scolor

        p2 = tf.add_paragraph()
        p2.text = "Key Activities:"
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = SLATE_DARK
        p2.space_before = Pt(8)

        p3 = tf.add_paragraph()
        p3.text = sacts
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = SLATE_DARK
        p3.space_before = Pt(4)

        p4 = tf.add_paragraph()
        p4.text = "Stage Outcome:"
        p4.font.size = Pt(10)
        p4.font.bold = True
        p4.font.color.rgb = SLATE_DARK
        p4.space_before = Pt(10)

        p5 = tf.add_paragraph()
        p5.text = f"Artifact:\n{sartifact}"
        p5.font.size = Pt(9.5)
        p5.font.bold = True
        p5.font.color.rgb = scolor
        p5.space_before = Pt(2)

    # ==========================================
    # SLIDE 6: IDEATION AND PROPOSED SOLUTION
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Ideation & Comparative Solution Rationale")
    add_footer(slide6, 6)

    # Comparison Table
    table_shape = slide6.shapes.add_table(5, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.3))
    table = table_shape.table

    # Column Widths
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(3.0)
    table.columns[2].width = Inches(3.0)
    table.columns[3].width = Inches(3.233)

    headers = ["Evaluation Metric", "Option A: Paper Register", "Option B: Google Form / Sheet", "CampusPlay Solution (Selected)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if i < 3 else BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    data = [
        ("Slot Visibility & Booking", "Zero visibility; requires visiting sports office in person", "Asynchronous; frequent overlapping submissions and conflicts", "Real-time visual slot grid; instant conflict prevention & booking"),
        ("Academic Timetable Integration", "Completely manual; impossible for staff to verify student schedules", "No validation; students self-certify without timetable checks", "Automated verification against branch/division semester timetable"),
        ("Equipment Inventory Governance", "Paper checkout book; uninspected returns and unaccounted damage", "Static spreadsheet; manual stock updates prone to delay", "Full lifecycle: check-and-set booking, return request, inspection modal"),
        ("Real-Time Synchronization", "None; telephone or notice board announcements", "None; email notifications with multi-hour delay", "Socket.io real-time room broadcasting for immediate updates")
    ]

    for row_idx, row_data in enumerate(data):
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if col_idx < 3 else LIGHT_GREEN
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(10)
            p.font.color.rgb = SLATE_DARK
            if col_idx == 0:
                p.font.bold = True
            elif col_idx == 3:
                p.font.bold = True
                p.font.color.rgb = RGBColor(6, 95, 70)

    # Bottom Takeaway Card
    add_card(slide6, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.95), bg_color=LIGHT_BLUE, border_color=BLUE)
    tbt = slide6.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.85))
    tft = tbt.text_frame
    tft.word_wrap = True
    pt = tft.paragraphs[0]
    pt.text = "SOLUTION RATIONALE & VALUE:"
    pt.font.size = Pt(10.5)
    pt.font.bold = True
    pt.font.color.rgb = BLUE
    pt2 = tft.add_paragraph()
    pt2.text = "CampusPlay transitions campus recreation from fragmented, error-prone manual intervention into an institutional-grade, verifiable service-oriented architecture with full governance, automated conflict prevention, and transparent auditing."
    pt2.font.size = Pt(11)
    pt2.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 7: SYSTEM ARCHITECTURE
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Technical Architecture: Multi-Tier Service Ecosystem")
    add_footer(slide7, 7)

    # 4 Architecture Tier Columns
    tiers = [
        ("PRESENTATION TIER", "Next.js 14 + Tailwind CSS",
         [("Client Engine", "Next.js 14 (Pages Router) + TypeScript"),
          ("State & Real-Time", "Socket.io-Client + Context Providers"),
          ("UI Components", "BookingWizard, GroundAvailability, ReturnModal"),
          ("Theme System", "Sport & Nature Harmony Palette + Lucide Icons")],
         BLUE),
        ("APPLICATION & API TIER", "Express.js RESTful Engine",
         [("HTTP Server", "Express 4.19 + Node.js (ES Modules)"),
          ("Security Layer", "Helmet, CORS Policy, Cookie-Parser"),
          ("Auth Guard", "JWT HttpOnly Cookies + Role RBAC Middleware"),
          ("Validation", "express-validator (Domain, Roll & Contact Checks)")],
         DEEP_BLUE),
        ("SERVICE BUSINESS LAYER", "Domain Logic & Automation",
         [("Timetable Engine", "Interval overlap & boundary conflict detection"),
          ("Equipment Service", "Check-and-set stock allocation & reconciliation"),
          ("Real-Time Broadcaster", "Socket.io room broadcasting (admin / student)"),
          ("Import / Storage", "CSV stream parser + binary magic-byte validator")],
         GREEN),
        ("PERSISTENCE TIER", "MongoDB Atlas Cloud Cluster",
         [("Data Store", "MongoDB Cloud Database (Atlas Replica Set)"),
          ("ODM Layer", "Mongoose 8.6 with Schemas & Compound Indexes"),
          ("Core Models", "User, Booking, Equipment, EquipmentBooking"),
          ("Audit Models", "Timetable, TimetableDocument, InventoryTransaction")],
         NAVY)
    ]

    for idx, (ttitle, tsub, titems, tcolor) in enumerate(tiers):
        left = Inches(0.8 + idx * 2.98)
        top = Inches(1.4)
        add_card(slide7, left, top, Inches(2.82), Inches(5.4), bg_color=WHITE, border_color=tcolor)

        # Header Box
        hb = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(2.82), Inches(0.7))
        hb.fill.solid()
        hb.fill.fore_color.rgb = tcolor
        hb.line.fill.background()
        htf = hb.text_frame
        htf.word_wrap = True
        hp = htf.paragraphs[0]
        hp.text = ttitle
        hp.font.size = Pt(11)
        hp.font.bold = True
        hp.font.color.rgb = WHITE
        hp.alignment = PP_ALIGN.CENTER

        hp2 = htf.add_paragraph()
        hp2.text = tsub
        hp2.font.size = Pt(9.5)
        hp2.font.color.rgb = LIGHT_GREEN if tcolor == GREEN else LIGHT_BLUE
        hp2.alignment = PP_ALIGN.CENTER

        # Content Box
        cb = slide7.shapes.add_textbox(left + Inches(0.12), top + Inches(0.8), Inches(2.58), Inches(4.5))
        ctf = cb.text_frame
        ctf.word_wrap = True

        for i_idx, (ititle, idesc) in enumerate(titems):
            p = ctf.add_paragraph() if i_idx > 0 else ctf.paragraphs[0]
            p.text = f"▶ {ititle}"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = tcolor
            p.space_before = Pt(8) if i_idx > 0 else Pt(2)

            p2 = ctf.add_paragraph()
            p2.text = idesc
            p2.font.size = Pt(9.5)
            p2.font.color.rgb = SLATE_DARK
            p2.space_before = Pt(1)

    # ==========================================
    # SLIDE 8: TECHNOLOGY STACK
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Verified Technology Stack & Rationale")
    add_footer(slide8, 8)

    tech_categories = [
        ("Frontend Technologies", BLUE, [
            ("Next.js 14.2 (TypeScript)", "Server-side rendering, type-safety, modular routing, and font optimization."),
            ("Tailwind CSS 3.4", "Custom utility-first styling with 'Sport & Nature Harmony' design tokens."),
            ("Socket.io-Client 4.8", "Bi-directional WebSocket client for zero-latency status synchronization."),
            ("Framer Motion & Lucide", "Hardware-accelerated micro-animations and clean sports UI iconography.")
        ]),
        ("Backend Technologies", DEEP_BLUE, [
            ("Node.js 18+ (ES Modules)", "Asynchronous, event-driven JavaScript runtime with native module syntax."),
            ("Express.js 4.19", "Lightweight, flexible web framework powering REST endpoints and middleware."),
            ("Socket.io 4.8 Server", "WebSocket server enabling real-time administrative and student room messaging."),
            ("Multer 2.4", "Multipart form handling for timetable CSV and reference document uploads.")
        ]),
        ("Data, Security & Testing", GREEN, [
            ("MongoDB Atlas + Mongoose 8.6", "Cloud NoSQL document store with strict Mongoose schema validation."),
            ("JWT + HttpOnly Cookies", "Tamper-proof authentication preventing XSS and credential interception."),
            ("Node Test Runner (node:test)", "Native zero-dependency test suite running 49 automated verification tests."),
            ("Render & Atlas Cloud", "Live production backend deployed on Render with cloud database peering.")
        ])
    ]

    for idx, (cat_title, cat_color, items) in enumerate(tech_categories):
        left = Inches(0.8 + idx * 3.98)
        top = Inches(1.4)
        add_card(slide8, left, top, Inches(3.78), Inches(5.4), bg_color=WHITE, border_color=cat_color)

        # Header Box
        hb = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(3.78), Inches(0.6))
        hb.fill.solid()
        hb.fill.fore_color.rgb = cat_color
        hb.line.fill.background()
        htf = hb.text_frame
        htf.word_wrap = True
        hp = htf.paragraphs[0]
        hp.text = cat_title.upper()
        hp.font.size = Pt(12)
        hp.font.bold = True
        hp.font.color.rgb = WHITE
        hp.alignment = PP_ALIGN.CENTER

        # Items
        cb = slide8.shapes.add_textbox(left + Inches(0.18), top + Inches(0.7), Inches(3.42), Inches(4.5))
        ctf = cb.text_frame
        ctf.word_wrap = True

        for i_idx, (t_name, t_rat) in enumerate(items):
            p = ctf.add_paragraph() if i_idx > 0 else ctf.paragraphs[0]
            p.text = f"✔ {t_name}"
            p.font.size = Pt(11.5)
            p.font.bold = True
            p.font.color.rgb = cat_color
            p.space_before = Pt(10) if i_idx > 0 else Pt(2)

            p2 = ctf.add_paragraph()
            p2.text = t_rat
            p2.font.size = Pt(10)
            p2.font.color.rgb = SLATE_DARK
            p2.space_before = Pt(2)

    # ==========================================
    # SLIDE 9: CORE FEATURES AND MODULES
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Core System Modules & Implemented Features")
    add_footer(slide9, 9)

    modules_data = [
        ("1. Multi-Step Booking Wizard", 
         "• Real-time visual ground availability grid across 11 daily slots (8 AM - 7 PM)\n• Support for single or dual consecutive hour reservations\n• Dynamic squad member registration with roll & email validation\n• Comprehensive pre-submission review modal before final booking creation",
         BLUE),
        ("2. Timetable Conflict Verification", 
         "• Compares requested slot against student's branch/division timetable\n• Detects overlaps, boundary touching (no collision), and multi-hour practicals\n• Flags bookings: CONFLICT DETECTED, NO CONFLICT, or DOCUMENT AVAILABLE\n• Provides sports administrators with objective academic cross-checks",
         RGBColor(239, 68, 68)),
        ("3. Equipment Inventory & Governance", 
         "• Catalog across Cricket, Football, Basketball, Volleyball, and Badminton\n• Check-and-set availability calculation preventing over-allocation\n• Student-initiated return workflow with timestamped return requests\n• Admin inspection modal reconciling returned, damaged, and missing stock",
         GREEN),
        ("4. Admin Operations Control Center", 
         "• 4-in-1 workspace: Facility Bookings, Equipment, Facilities, Timetables\n• Status-filtered queues (Pending, Approved, Rejected, Cancelled)\n• One-click Approve / Reject with mandatory rejection reason logging\n• Full decision revocation capability (moving bookings back to pending)",
         DEEP_BLUE)
    ]

    for idx, (mtitle, mdesc, mcolor) in enumerate(modules_data):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.35 + row * 2.8)
        add_card(slide9, left, top, Inches(5.75), Inches(2.6), bg_color=WHITE, border_color=mcolor)

        tb = slide9.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.35), Inches(2.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = mtitle
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = mcolor

        p2 = tf.add_paragraph()
        p2.text = mdesc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = SLATE_DARK
        p2.space_before = Pt(6)

    # ==========================================
    # SLIDE 10: USER JOURNEY AND WORKFLOW
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "End-to-End User Journeys: Student & Admin Workflows")
    add_footer(slide10, 10)

    # Top Flow: Student Ground Booking Journey
    add_card(slide10, Inches(0.8), Inches(1.35), Inches(11.733), Inches(2.65), bg_color=WHITE, border_color=BLUE)
    tb_flow1 = slide10.shapes.add_textbox(Inches(1.0), Inches(1.45), Inches(11.3), Inches(2.45))
    tf1 = tb_flow1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "STUDENT GROUND RESERVATION WORKFLOW"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE

    steps1 = [
        ("Step 1: Auth & Profile", "Login with @vit.edu.in; system validates roll & fetches branch/division/semester profile."),
        ("Step 2: Facility & Slot", "Select facility (Ground 1/2, Court); pick date; select vacant 1-2 hour time-slot."),
        ("Step 3: Squad & Gear", "Add teammate roll/emails; optionally request sports equipment (footballs, bats)."),
        ("Step 4: Review & Submit", "Confirm all details in review modal; submission locks slot & generates Pending request."),
        ("Step 5: Live Notification", "Socket.io notifies student instantly when Admin reviews request with reason if rejected.")
    ]
    for s_title, s_desc in steps1:
        p = tf1.add_paragraph()
        p.text = f"▶ {s_title} ➔ {s_desc}"
        p.font.size = Pt(10)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(3)

    # Bottom Flow: Equipment Issuance & Return Inspection Workflow
    add_card(slide10, Inches(0.8), Inches(4.2), Inches(11.733), Inches(2.7), bg_color=WHITE, border_color=GREEN)
    tb_flow2 = slide10.shapes.add_textbox(Inches(1.0), Inches(4.3), Inches(11.3), Inches(2.5))
    tf2 = tb_flow2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "EQUIPMENT ISSUANCE & RETURN INSPECTION WORKFLOW"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN

    steps2 = [
        ("Phase 1: Allocation", "Student requests equipment with ground booking; system checks check-and-set stock."),
        ("Phase 2: Admin Approval", "Admin reviews request; marks status as 'Approved' / 'Ready for Collection'."),
        ("Phase 3: Issuance", "Student collects gear at sports desk; Admin clicks 'Mark Issued' setting timestamp & issuer."),
        ("Phase 4: Return Request", "Student returns gear; clicks 'Submit Return' updating status to 'Return Requested'."),
        ("Phase 5: Reconciliation", "Admin inspects gear: confirms usable stock, logs damaged/missing units, and commits audit trail.")
    ]
    for s_title, s_desc in steps2:
        p = tf2.add_paragraph()
        p.text = f"▶ {s_title} ➔ {s_desc}"
        p.font.size = Pt(10)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(3)

    # ==========================================
    # SLIDE 11: UI/UX AND PROTOTYPE
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "UI/UX Design System: Sport & Nature Harmony")
    add_footer(slide11, 11)

    ui_cards = [
        ("Sport & Nature Harmony Palette", BLUE, [
            ("Primary Blue (#2563EB)", "High-contrast action color for primary CTAs, active tabs, and navigation links."),
            ("Emerald Green (#10B981)", "Represents campus sports turf, vitality, and successful approval badges."),
            ("Slate Dark (#0F172A)", "Deep navy-slate background ensuring high visual hierarchy and readability."),
            ("Status Colors", "Green (Approved), Amber (Pending), Red (Rejected), Gray (Cancelled).")
        ]),
        ("Typography & Micro-Interactions", GREEN, [
            ("Dual Font System", "Poppins (600/700) for impactful headings + Inter (400/500) for body legibility."),
            ("Framer Motion Transitions", "Smooth 200ms ease-out transitions on modals, cards, and buttons."),
            ("Glassmorphism & Shadows", "Subtle multi-layered neumorphic shadows and border highlights."),
            ("Responsive Viewport-First", "Fluid layout optimized across mobile smartphones, tablets, and desktop.")
        ]),
        ("Usability & Form Resilience", DEEP_BLUE, [
            ("Contextual Guidance", "Inline validation errors, roll number format enforcement, and domain lock."),
            ("Pre-Submission Review", "Confirmation modal displaying all booking details prior to commit."),
            ("Safe Empty States", "Descriptive empty state illustrations when no bookings or requests exist."),
            ("Zero Silent Overwrites", "Verified profile data protected against student tampering.")
        ])
    ]

    for idx, (utitle, ucolor, uitems) in enumerate(ui_cards):
        left = Inches(0.8 + idx * 3.98)
        top = Inches(1.4)
        add_card(slide11, left, top, Inches(3.78), Inches(5.4), bg_color=WHITE, border_color=ucolor)

        hb = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(3.78), Inches(0.6))
        hb.fill.solid()
        hb.fill.fore_color.rgb = ucolor
        hb.line.fill.background()
        htf = hb.text_frame
        htf.word_wrap = True
        hp = htf.paragraphs[0]
        hp.text = utitle.upper()
        hp.font.size = Pt(11.5)
        hp.font.bold = True
        hp.font.color.rgb = WHITE
        hp.alignment = PP_ALIGN.CENTER

        cb = slide11.shapes.add_textbox(left + Inches(0.18), top + Inches(0.7), Inches(3.42), Inches(4.5))
        ctf = cb.text_frame
        ctf.word_wrap = True

        for i_idx, (t_name, t_desc) in enumerate(uitems):
            p = ctf.add_paragraph() if i_idx > 0 else ctf.paragraphs[0]
            p.text = f"• {t_name}"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = ucolor
            p.space_before = Pt(8) if i_idx > 0 else Pt(2)

            p2 = ctf.add_paragraph()
            p2.text = t_desc
            p2.font.size = Pt(9.5)
            p2.font.color.rgb = SLATE_DARK
            p2.space_before = Pt(2)

    # ==========================================
    # SLIDE 12: IMPLEMENTATION AND DEMONSTRATION
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "Technical Implementation & Live Demonstration Sequence")
    add_footer(slide12, 12)

    # Left Column: Implementation Highlights
    add_card(slide12, Inches(0.8), Inches(1.4), Inches(5.75), Inches(5.4))
    b1 = slide12.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(5.35), Inches(5.1))
    t1 = b1.text_frame
    t1.word_wrap = True

    p = t1.paragraphs[0]
    p.text = "⚙️ Core Implementation Architecture"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = BLUE

    impl_points = [
        ("Service-Layer Decoupling", "Controllers remain slim; business logic lives in dedicated services (timetableConflictService, equipmentService, socketService)."),
        ("Interval Overlap Mathematics", "Custom time-parser converts '10–11 AM' into normalized 24-hr intervals (10:00-11:00); mathematically evaluates [startA < endB && endA > startB] to detect true overlaps while permitting touching boundaries."),
        ("Check-and-Set Inventory Idempotency", "Atomic inventory checks prevent race conditions and duplicate restorations during concurrent return requests."),
        ("Binary Magic-Byte Inspection", "Uploaded timetable reference files are checked against raw binary byte headers (%PDF-, PNG, JPEG, WebP) to prevent extension spoofing and malware injection.")
    ]
    for ititle, idesc in impl_points:
        p = t1.add_paragraph()
        p.text = f"▶ {ititle}: {idesc}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(8)

    # Right Column: Live Demonstration Flow
    add_card(slide12, Inches(6.78), Inches(1.4), Inches(5.75), Inches(5.4))
    b2 = slide12.shapes.add_textbox(Inches(6.98), Inches(1.55), Inches(5.35), Inches(5.1))
    t2 = b2.text_frame
    t2.word_wrap = True

    p = t2.paragraphs[0]
    p.text = "🎬 Step-by-Step Live Demonstration Plan"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = GREEN

    demo_steps = [
        ("Step 1: Student Login & Booking", "Log in as student (roll: 24IT1C0001); navigate to /bookings; observe pre-filled academic profile (INFT, SE, Div A)."),
        ("Step 2: Ground & Slot Selection", "Choose Football on Ground 1; observe green/gray availability grid; select 10–11 AM slot and 2 footballs."),
        ("Step 3: Review & Submit", "Inspect review modal with all squad details; submit booking; verify slot switches to Reserved."),
        ("Step 4: Admin Inspection", "Open Admin Dashboard in second window; observe instant real-time notification; review Academic Verification report showing 'CONFLICT DETECTED'."),
        ("Step 5: Equipment Reconciliation", "Navigate to Equipment tab; issue gear; trigger student return request; confirm return with damaged stock audit.")
    ]
    for dtitle, ddesc in demo_steps:
        p = t2.add_paragraph()
        p.text = f"• {dtitle}: {ddesc}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 13: TESTING AND VERIFIED RESULTS
    # ==========================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "Automated Testing & Empirical Verification")
    add_footer(slide13, 13)

    # Top Banner: 49/49 Tests Passed
    add_card(slide13, Inches(0.8), Inches(1.35), Inches(11.733), Inches(1.1), bg_color=LIGHT_GREEN, border_color=GREEN)
    tb_pass = slide13.shapes.add_textbox(Inches(1.0), Inches(1.42), Inches(11.3), Inches(0.95))
    tfp = tb_pass.text_frame
    tfp.word_wrap = True
    pp = tfp.paragraphs[0]
    pp.text = "EMPIRICAL VERIFICATION: 49 AUTOMATED TESTS EXECUTED • 49 PASSED (100% SUCCESS RATE)"
    pp.font.size = Pt(13)
    pp.font.bold = True
    pp.font.color.rgb = RGBColor(6, 95, 70)
    pp2 = tfp.add_paragraph()
    pp2.text = "Executed using Node.js native test runner (node:test, node:assert/strict). All tests pass with zero mock failures, verifying backend stability, data integrity, security bounds, and business logic."
    pp2.font.size = Pt(10.5)
    pp2.font.color.rgb = SLATE_DARK

    # 3 Test Suite Breakdown Cards
    suites = [
        ("Suite 1: Equipment Governance (10 Tests)", BLUE, [
            ("Full Return Restores Usable Stock", "Verified inventory increases by exact usable returned units."),
            ("Partial Return & Damage Reconciliation", "Damaged and missing items logged without inflating usable stock."),
            ("Check-and-Set Idempotency", "Duplicate confirmation requests rejected; stock not restored twice."),
            ("Unauthorized Student Block", "Students barred from submitting returns for other peers."),
            ("Cancellation Integrity", "Cancelled requests release allocations without inventory leak.")
        ]),
        ("Suite 2: Academic Profile & Contact (15 Tests)", GREEN, [
            ("Strict Phone Validation", "Mandatory 10-digit primary phone; optional distinct alternate."),
            ("Profile Overwrite Immunity", "Verified academic profile cannot be silently modified by client."),
            ("Academic Combination Bounds", "Invalid stream/year/division/semester combinations rejected."),
            ("Data Privacy & RBAC", "Contact numbers restricted strictly to authorized sports admins."),
            ("Safe Defaults for Existing Users", "Legacy records handled gracefully without server exceptions.")
        ]),
        ("Suite 3: Conflict & Document Security (24 Tests)", DEEP_BLUE, [
            ("Slot Interval Overlap Detection", "True overlaps flagged; exact boundary-touching slots allowed."),
            ("Division & Branch Segregation", "Correct division timetable queried per student record."),
            ("Multi-Hour Consecutive Checking", "Full multi-slot reservation tested against academic schedule."),
            ("Binary Magic-Byte Inspection", "Renamed text/spoofed files rejected; %PDF- & image headers verified."),
            ("Anti-Path Traversal", "Malicious storage keys sanitized preventing directory escape.")
        ])
    ]

    for idx, (stitle, scolor, sitems) in enumerate(suites):
        left = Inches(0.8 + idx * 3.98)
        top = Inches(2.6)
        add_card(slide13, left, top, Inches(3.78), Inches(4.3), bg_color=WHITE, border_color=scolor)

        hb = slide13.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(3.78), Inches(0.55))
        hb.fill.solid()
        hb.fill.fore_color.rgb = scolor
        hb.line.fill.background()
        htf = hb.text_frame
        htf.word_wrap = True
        hp = htf.paragraphs[0]
        hp.text = stitle
        hp.font.size = Pt(10.5)
        hp.font.bold = True
        hp.font.color.rgb = WHITE
        hp.alignment = PP_ALIGN.CENTER

        cb = slide13.shapes.add_textbox(left + Inches(0.15), top + Inches(0.62), Inches(3.48), Inches(3.5))
        ctf = cb.text_frame
        ctf.word_wrap = True

        for i_idx, (t_name, t_desc) in enumerate(sitems):
            p = ctf.add_paragraph() if i_idx > 0 else ctf.paragraphs[0]
            p.text = f"✔ {t_name}"
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = scolor
            p.space_before = Pt(5) if i_idx > 0 else Pt(2)

            p2 = ctf.add_paragraph()
            p2.text = t_desc
            p2.font.size = Pt(9)
            p2.font.color.rgb = SLATE_DARK
            p2.space_before = Pt(1)

    # ==========================================
    # SLIDE 14: CHALLENGES, LIMITATIONS, AND FUTURE SCOPE
    # ==========================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_header(slide14, "Challenges, Current Limitations & Strategic Roadmap")
    add_footer(slide14, 14)

    # 3 Columns
    # Col 1: Technical Challenges Solved
    add_card(slide14, Inches(0.8), Inches(1.4), Inches(3.64), Inches(5.4))
    b1 = slide14.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(3.24), Inches(5.1))
    t1 = b1.text_frame
    t1.word_wrap = True
    p = t1.paragraphs[0]
    p.text = "🛠 Technical Challenges Overcome"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE

    challenges = [
        ("Boundary-Touching Slots", "Solved false positive conflicts when a booking starts exactly when a class ends by normalizing intervals into [start < endSession && end > startSession]."),
        ("Equipment Race Conditions", "Prevented concurrent double-returns by implementing check-and-set idempotency."),
        ("File Upload Security", "Prevented spoofed file uploads by parsing magic bytes rather than trusting file extensions.")
    ]
    for ctitle, cdesc in challenges:
        p = t1.add_paragraph()
        p.text = f"• {ctitle}: {cdesc}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(10)

    # Col 2: Current Limitations
    add_card(slide14, Inches(4.84), Inches(1.4), Inches(3.64), Inches(5.4))
    b2 = slide14.shapes.add_textbox(Inches(5.04), Inches(1.55), Inches(3.24), Inches(5.1))
    t2 = b2.text_frame
    t2.word_wrap = True
    p = t2.paragraphs[0]
    p.text = "⚠️ Known Limitations"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = AMBER

    limits = [
        ("Physical Access Enforcement", "System manages digital permissions; physical ground turnstiles or gate locks require human check-in."),
        ("Reference Document OCR", "Scanned PDF/image timetables require visual admin inspection rather than automated OCR text extraction."),
        ("Single Institution Scope", "Rosters and authentication currently configured specifically for Vidyalankar Institute of Technology (@vit.edu.in).")
    ]
    for ltitle, ldesc in limits:
        p = t2.add_paragraph()
        p.text = f"• {ltitle}: {ldesc}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(10)

    # Col 3: Future Scope & Enhancements
    add_card(slide14, Inches(8.88), Inches(1.4), Inches(3.64), Inches(5.4))
    b3 = slide14.shapes.add_textbox(Inches(9.08), Inches(1.55), Inches(3.24), Inches(5.1))
    t3 = b3.text_frame
    t3.word_wrap = True
    p = t3.paragraphs[0]
    p.text = "🚀 Future Strategic Roadmap"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GREEN

    future = [
        ("Smart RFID / QR Turnstiles", "Automated physical gate unlock upon scanning booking QR code on student mobile."),
        ("Inter-College Tournaments", "Dedicated fixture scheduling, league tables, and team registration for annual campus fests."),
        ("Automated Equipment Fines", "Integration with campus fee portal to automatically assess replacement fines for unreturned or damaged equipment.")
    ]
    for ftitle, fdesc in future:
        p = t3.add_paragraph()
        p.text = f"• {ftitle}: {fdesc}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = SLATE_DARK
        p.space_before = Pt(10)

    # ==========================================
    # SLIDE 15: CONCLUSION AND Q&A
    # ==========================================
    slide15 = prs.slides.add_slide(blank_layout)
    bg15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg15.fill.solid()
    bg15.fill.fore_color.rgb = NAVY
    bg15.line.fill.background()

    # Title Box
    tbox15 = slide15.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.5))
    tf15 = tbox15.text_frame
    tf15.word_wrap = True

    p = tf15.paragraphs[0]
    p.text = "PROJECT SUMMARY & CONCLUSION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN

    p2 = tf15.add_paragraph()
    p2.text = "Transforming Campus Athletics Through Engineered Design Thinking"
    p2.font.size = Pt(30)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(6)

    # 3 Summary Cards
    scards = [
        ("Problem Solved", "Replaced opaque, manual, and conflict-prone paper ground scheduling with a transparent, real-time sports reservation platform.", BLUE),
        ("Design Thinking Impact", "Applied human-centered methodology (Empathize to Test), addressing genuine collegiate student and faculty pain points.", GREEN),
        ("Engineering Rigor", "Built full-stack Next.js + Express + MongoDB architecture with Socket.io live sync and 49 verified automated unit/integration tests.", DEEP_BLUE)
    ]

    for idx, (stitle, sdesc, scolor) in enumerate(scards):
        left = Inches(1.0 + idx * 3.85)
        top = Inches(2.6)
        card = add_card(slide15, left, top, Inches(3.6), Inches(2.5), bg_color=RGBColor(30, 41, 59), border_color=scolor)

        tb = slide15.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(3.2), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = scolor

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(11)
        p2.font.color.rgb = LIGHT_BLUE
        p2.space_before = Pt(10)

    # Q&A Prompt Box
    add_card(slide15, Inches(1.0), Inches(5.4), Inches(11.333), Inches(1.5), bg_color=RGBColor(24, 34, 53), border_color=BLUE)
    tb_qa = slide15.shapes.add_textbox(Inches(1.2), Inches(5.5), Inches(10.9), Inches(1.3))
    tfqa = tb_qa.text_frame
    tfqa.word_wrap = True

    p = tfqa.paragraphs[0]
    p.text = "THANK YOU! QUESTIONS & ANSWERS (VIVA DEFENSE)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    p2 = tfqa.add_paragraph()
    p2.text = "Aryan Kelkar & Project Team | Vidyalankar Institute of Technology (VIT) | Department of Information Technology\nRepository: https://github.com/aryankelkar/CampusPLAY | Backend Live: https://campusplay-4g0z.onrender.com"
    p2.font.size = Pt(11)
    p2.font.color.rgb = GREEN
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(8)

    prs.save(output_path)
    print(f"[SUCCESS] Successfully created 15-slide PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\CampusPLAY\deliverables"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "CampusPlay_Presentation.pptx")
    create_presentation(out_file)
